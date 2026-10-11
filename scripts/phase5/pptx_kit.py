"""Small python-pptx toolkit for an editable, vector-only institutional deck.

Every paragraph receives an explicit ``a:endParaRPr``: PowerPoint's PDF exporter omits
hyperlink annotations from paragraphs without one, which python-pptx does not add.
"""

from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---- design tokens -----------------------------------------------------------------
NAVY = "0B2341"
INK = "1F2933"
MUTED = "4F5B6A"
FAINT = "8A94A3"
RULE = "D3DAE3"
PANEL = "F1F4F8"
PANEL_DARK = "E3E9F1"
WHITE = "FFFFFF"
BLUE_250 = "86B6EF"
BLUE_450 = "2A78D6"
BLUE_550 = "1C5CAB"
BLUE_700 = "0D366B"
GRAY_BAR = "9AA5B1"
RED = "B23A3A"
AMBER = "9A6A12"
LINK = "1F5FA8"
HEAD_FONT = "Cambria"
BODY_FONT = "Arial"
CODE_FONT = "Consolas"

SLIDE_W = 13.333
SLIDE_H = 7.5
LEFT = 0.6
RIGHT = SLIDE_W - 0.6
WIDTH = RIGHT - LEFT


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def inch(value: float) -> Emu:
    return Inches(value)


# ---- theme -------------------------------------------------------------------------
def apply_theme(prs, *, heading=HEAD_FONT, body=BODY_FONT, link=LINK) -> None:
    """Set theme fonts and hyperlink colors so nothing falls back to Calibri."""
    theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    xml = theme_part.blob.decode("utf-8")
    xml = re.sub(
        r'(<a:majorFont>\s*<a:latin typeface=")[^"]*(")', rf"\g<1>{heading}\g<2>", xml, count=1
    )
    xml = re.sub(
        r'(<a:minorFont>\s*<a:latin typeface=")[^"]*(")', rf"\g<1>{body}\g<2>", xml, count=1
    )
    for tag in ("hlink", "folHlink"):
        xml = re.sub(
            rf"(<a:{tag}>\s*<a:srgbClr val=\")[0-9A-Fa-f]{{6}}(\")", rf"\g<1>{link}\g<2>", xml
        )
    xml = re.sub(r'name="Office Theme"', 'name="Everpure Institutional"', xml, count=1)
    theme_part._blob = xml.encode("utf-8")


# ---- text ----------------------------------------------------------------------------
@dataclass
class Run:
    text: str
    bold: bool | None = None
    italic: bool | None = None
    color: str | None = None
    size: float | None = None
    font: str | None = None
    link: str | None = None
    underline: bool | None = None
    superscript: bool = False


@dataclass
class Para:
    runs: list[Run] | str
    size: float = 12
    bold: bool = False
    color: str = INK
    font: str = BODY_FONT
    align: str = "l"
    bullet: str | None = None
    indent: float = 0.0
    space_before: float = 0
    space_after: float = 4
    line_spacing: float | None = None
    bullet_color: str | None = None
    level_indent: float = 0.0
    extra: dict = field(default_factory=dict)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}


def _end_para(p, size: float, font: str) -> None:
    for old in p._p.findall(qn("a:endParaRPr")):
        p._p.remove(old)
    end = etree.SubElement(p._p, qn("a:endParaRPr"))
    end.set("lang", "en-US")
    end.set("sz", str(int(round(size * 100))))
    end.set("dirty", "0")
    latin = etree.SubElement(end, qn("a:latin"))
    latin.set("typeface", font)


def _set_bullet(p, char: str, indent: float, color: str, level_indent: float) -> None:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent + level_indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone", "a:buAutoNum"):
        for old in pPr.findall(qn(tag)):
            pPr.remove(old)
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    font = etree.SubElement(pPr, qn("a:buFont"))
    font.set("typeface", BODY_FONT)
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def _no_bullet(p) -> None:
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("a:buNone")) is None:
        etree.SubElement(pPr, qn("a:buNone"))


def fill_text_frame(tf, paras: list[Para]) -> None:
    for index, spec in enumerate(paras):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = ALIGN[spec.align]
        p.space_before = Pt(spec.space_before)
        p.space_after = Pt(spec.space_after)
        if spec.line_spacing:
            p.line_spacing = spec.line_spacing
        if spec.bullet:
            _set_bullet(
                p,
                spec.bullet,
                spec.indent or 0.18,
                spec.bullet_color or spec.color,
                spec.level_indent,
            )
        else:
            _no_bullet(p)
        runs = [Run(spec.runs)] if isinstance(spec.runs, str) else spec.runs
        for item in runs:
            r = p.add_run()
            r.text = item.text
            f = r.font
            f.name = item.font or spec.font
            f.size = Pt(item.size or spec.size)
            f.bold = spec.bold if item.bold is None else item.bold
            if item.italic is not None:
                f.italic = item.italic
            f.color.rgb = rgb(item.color or spec.color)
            if item.link:
                r.hyperlink.address = item.link
                f.color.rgb = rgb(LINK)
                f.underline = True if item.underline is None else item.underline
            r._r.get_or_add_rPr().set("lang", "en-US")
            if item.superscript:
                r._r.get_or_add_rPr().set("baseline", "30000")
        _end_para(p, spec.size, spec.font)


def text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    paras: list[Para] | Para | str,
    *,
    anchor: str = "t",
    margin: tuple[float, float, float, float] = (0, 0, 0, 0),
    name: str | None = None,
    wrap: bool = True,
):
    if isinstance(paras, str):
        paras = [Para(paras)]
    elif isinstance(paras, Para):
        paras = [paras]
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (
        inch(margin[0]),
        inch(margin[1]),
        inch(margin[2]),
        inch(margin[3]),
    )
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[
        anchor
    ]
    fill_text_frame(tf, paras)
    return box


# ---- shapes ----------------------------------------------------------------------------
def rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str | None,
    *,
    line: str | None = None,
    line_w: float = 0.75,
    shape=MSO_SHAPE.RECTANGLE,
    name: str | None = None,
    radius: float | None = None,
):
    s = slide.shapes.add_shape(shape, inch(x), inch(y), inch(w), inch(h))
    if name:
        s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return s


def shape_text(
    s,
    paras: list[Para] | Para | str,
    *,
    anchor: str = "m",
    margin: tuple[float, float, float, float] = (0.08, 0.04, 0.08, 0.04),
) -> None:
    if isinstance(paras, str):
        paras = [Para(paras)]
    elif isinstance(paras, Para):
        paras = [paras]
    tf = s.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (
        inch(margin[0]),
        inch(margin[1]),
        inch(margin[2]),
        inch(margin[3]),
    )
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[
        anchor
    ]
    fill_text_frame(tf, paras)


def arrow(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = MUTED,
    width: float = 1.25,
    head: bool = True,
    dash: bool = False,
):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if head:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


# ---- tables -----------------------------------------------------------------------------
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _cell_props(cell, fill: str | None, borders: dict[str, tuple[str, float] | None]) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:noFill", "a:solidFill"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    position = 0
    for side in ("L", "R", "T", "B"):
        spec = borders.get(side)
        ln = etree.Element(qn(f"a:ln{side}"))
        if spec is None:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        else:
            color, width = spec
            ln.set("w", str(int(width * 12700)))
            ln.set("cap", "flat")
            ln.set("cmpd", "sng")
            ln.set("algn", "ctr")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
            etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        tcPr.insert(position, ln)
        position += 1
    if fill is None:
        tcPr.insert(position, etree.Element(qn("a:noFill")))
    else:
        sf = etree.Element(qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", fill)
        tcPr.insert(position, sf)


@dataclass
class Cell:
    value: str | list[Run]
    bold: bool | None = None
    color: str | None = None
    fill: str | None = None
    align: str | None = None
    size: float | None = None
    font: str | None = None


def table(
    slide,
    x: float,
    y: float,
    col_widths: list[float],
    rows: list[list[str | Cell]],
    *,
    row_h: float = 0.3,
    header_rows: int = 1,
    size: float = 11,
    header_size: float | None = None,
    header_fill: str | None = NAVY,
    header_color: str = WHITE,
    zebra: str | None = None,
    aligns: list[str] | None = None,
    rule: str = RULE,
    total_row: bool = False,
    margins: tuple[float, float] = (0.06, 0.03),
    name: str | None = None,
    row_heights: list[float] | None = None,
):
    n_rows, n_cols = len(rows), len(col_widths)
    gf = slide.shapes.add_table(
        n_rows, n_cols, inch(x), inch(y), inch(sum(col_widths)), inch(row_h * n_rows)
    )
    if name:
        gf.name = name
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    tbl_pr = gf._element.graphic.graphicData.tbl.tblPr
    style = tbl_pr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tbl_pr, qn("a:tableStyleId"))
    style.text = NO_STYLE_TABLE
    for j, w in enumerate(col_widths):
        tbl.columns[j].width = inch(w)
    for i in range(n_rows):
        tbl.rows[i].height = inch(row_heights[i] if row_heights else row_h)
    for i, row in enumerate(rows):
        is_header = i < header_rows
        is_total = total_row and i == n_rows - 1
        for j in range(n_cols):
            raw = row[j] if j < len(row) else ""
            spec = raw if isinstance(raw, Cell) else Cell(raw)
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = inch(margins[0])
            cell.margin_top = cell.margin_bottom = inch(margins[1])
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill = spec.fill
            if fill is None:
                if is_header:
                    fill = header_fill
                elif zebra and (i - header_rows) % 2 == 1:
                    fill = zebra
            borders: dict[str, tuple[str, float] | None] = {"L": None, "R": None, "T": None}
            borders["B"] = (rule, 0.75) if not is_header else None
            if is_total:
                borders["T"] = (INK, 1.0)
            _cell_props(cell, fill, borders)
            align = spec.align or (aligns[j] if aligns else "l")
            color = spec.color or (header_color if is_header else INK)
            bold = spec.bold if spec.bold is not None else (is_header or is_total)
            cell_size = spec.size or ((header_size or size) if is_header else size)
            runs = spec.value if isinstance(spec.value, list) else [Run(str(spec.value))]
            tf = cell.text_frame
            tf.word_wrap = True
            fill_text_frame(
                tf,
                [
                    Para(
                        runs,
                        size=cell_size,
                        bold=bold,
                        color=color,
                        font=spec.font or BODY_FONT,
                        align=align,
                        space_after=0,
                    )
                ],
            )
    return gf


# ---- charts ------------------------------------------------------------------------------
def style_chart_text(chart, size: float = 11, color: str = INK) -> None:
    chart.font.name = BODY_FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(color)


def bar_chart(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    categories: list[str],
    series: list[tuple[str, list[float]]],
    *,
    chart_type,
    colors: list[str],
    point_colors: list[str] | None = None,
    number_format: str = "#,##0",
    label_format: str | None = None,
    legend: bool = True,
    gap_width: int = 60,
    overlap: int = 0,
    font_size: float = 11,
    label_size: float | None = None,
    show_value_axis: bool = False,
    value_axis_min: float | None = None,
    value_axis_max: float | None = None,
    label_position=XL_LABEL_POSITION.OUTSIDE_END,
    gridlines: bool = False,
    name: str | None = None,
    category_axis_line: bool = True,
    reverse_categories: bool = False,
):
    data = CategoryChartData()
    data.categories = categories
    for label, values in series:
        data.add_series(label, values, number_format=number_format)
    gf = slide.shapes.add_chart(chart_type, inch(x), inch(y), inch(w), inch(h), data)
    if name:
        gf.name = name
    chart = gf.chart
    style_chart_text(chart, font_size)
    chart.has_title = False
    chart.has_legend = legend and len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(font_size)
        chart.legend.font.color.rgb = rgb(MUTED)
    plot = chart.plots[0]
    plot.gap_width = gap_width
    plot.overlap = overlap
    plot.vary_by_categories = False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = label_format or number_format
    labels.number_format_is_linked = False
    labels.position = label_position
    labels.font.size = Pt(label_size or font_size)
    labels.font.color.rgb = rgb(INK)
    for s_index, s in enumerate(plot.series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(colors[s_index % len(colors)])
        s.format.line.fill.background()
        s.invert_if_negative = False
    if point_colors:
        for p_index, color in enumerate(point_colors):
            point = plot.series[0].points[p_index]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb(color)
    va = chart.value_axis
    va.has_major_gridlines = gridlines
    if gridlines:
        va.major_gridlines.format.line.color.rgb = rgb(RULE)
        va.major_gridlines.format.line.width = Pt(0.5)
    va.visible = show_value_axis
    va.tick_labels.font.size = Pt(font_size)
    va.tick_labels.font.color.rgb = rgb(MUTED)
    va.format.line.fill.background()
    if value_axis_min is not None:
        va.minimum_scale = value_axis_min
    if value_axis_max is not None:
        va.maximum_scale = value_axis_max
    ca = chart.category_axis
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(font_size)
    ca.tick_labels.font.color.rgb = rgb(INK)
    ca.has_major_gridlines = False
    if category_axis_line:
        ca.format.line.color.rgb = rgb(RULE)
        ca.format.line.width = Pt(0.75)
    else:
        ca.format.line.fill.background()
    if reverse_categories:
        ca.reverse_order = True
    return gf


def clone(element):
    return copy.deepcopy(element)
